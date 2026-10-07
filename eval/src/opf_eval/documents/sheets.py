"""Spreadsheet parsers: XLSX (openpyxl) and CSV/TSV (stdlib `csv`).

Both use the same layout. Each non-empty row becomes one line with its
non-empty cells joined by a tab, and each non-empty cell is its own segment
so a detected span maps straight back to a cell:

    XLSX  {"sheet": "People", "cell": "B2"}     sheets separated by a blank line
    CSV   {"row": 2, "column": "email"}         row 1 is the header row

Empty cells are left out of the line rather than kept as runs of tabs. The
segment map already records each value's column, and long runs of tabs waste
detector context.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
from pathlib import Path

from .base import TextBuilder
from .text import decode_text

FILES_HINT = "pip install 'opf-eval[files]'"


def render_value(value) -> str:
    """Render a cell value the way a person reading the sheet would see it."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer() and abs(value) < 1e16:
        # Phone numbers and IDs stored as numbers would otherwise read
        # "5551234567.0", which no detector recognises.
        return str(int(value))
    if isinstance(value, dt.datetime) and value.time() == dt.time() and value.tzinfo is None:
        # Excel stores dates as datetimes; a birth date should not read
        # "1990-05-17T00:00:00".
        return value.date().isoformat()
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, dt.timedelta):
        return str(value)
    return str(value)


def column_letter(index: int) -> str:
    """1 -> "A", 27 -> "AA" (same as openpyxl.utils.get_column_letter)."""
    letters = ""
    while index > 0:
        index, rem = divmod(index - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def extract_xlsx(path: Path) -> TextBuilder:
    try:
        import openpyxl
    except ImportError as e:  # pragma: no cover - depends on the environment
        raise ImportError(f"Parsing .xlsx files needs openpyxl: {FILES_HINT}") from e

    b = TextBuilder()
    wb = openpyxl.load_workbook(Path(path), read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            first_row = True
            # Read-only mode trusts the sheet's stored <dimension> to bound
            # iter_rows, and some non-Excel writers store a wrong one ("A1"),
            # which would silently drop every other cell. Forget it so rows
            # and columns run as far as the cells really go. The fixed origin
            # keeps cell references right whatever the dimension said.
            ws.reset_dimensions()
            for r, row in enumerate(ws.iter_rows(min_row=1, min_col=1, values_only=True), 1):
                cells = [(c, render_value(v)) for c, v in enumerate(row, 1)]
                cells = [(c, v) for c, v in cells if v.strip()]
                if not cells:
                    continue
                if first_row:
                    b.blank_line()
                    first_row = False
                else:
                    b.newline()
                for k, (c, value) in enumerate(cells):
                    if k:
                        b.sep("\t")
                    b.add(value, {"sheet": ws.title, "cell": f"{column_letter(c)}{r}"})
    finally:
        wb.close()
    return b


def _sniff_dialect(sample: str):
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        return csv.excel


def extract_csv(path: Path) -> TextBuilder:
    path = Path(path)
    text, warnings = decode_text(path.read_bytes())
    if path.suffix.lower() == ".tsv":
        dialect = csv.excel_tab
    else:
        dialect = _sniff_dialect(text[:64 * 1024])
    reader = csv.reader(io.StringIO(text, newline=""), dialect)

    b = TextBuilder()
    b.warnings.extend(warnings)
    header: list[str] = []
    for r, row in enumerate(reader, 1):
        if r == 1:
            header = [h.strip() for h in row]
        cells = [(c, v) for c, v in enumerate(row, 1) if v.strip()]
        if not cells:
            continue
        b.newline()
        for k, (c, value) in enumerate(cells):
            if k:
                b.sep("\t")
            name = header[c - 1] if c <= len(header) and header[c - 1] else column_letter(c)
            b.add(value, {"row": r, "column": name})
    return b
