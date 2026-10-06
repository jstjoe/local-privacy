"""File parsing (`opf_eval.documents`): every parser keeps an exact offset map
from the extracted text back to page, cell, line or header.

Sample files are built in a session fixture with the same libraries the
parsers read them with. The one PDF is committed under tests/data/documents
because no PDF writer with a text layer is a dependency.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import itertools
import subprocess
import sys
from email.message import EmailMessage
from pathlib import Path

import pytest

from opf_eval import documents
from opf_eval.documents import Document, Segment, chunk, from_text, parse, parse_many
from opf_eval.documents.base import TextBuilder
from opf_eval.documents.html import html_to_lines

DATA = Path(__file__).parent / "data" / "documents"
EMAIL = "jane.doe@example.com"

CSV_TEXT = "name,email,phone\r\nJane Doe,jane.doe@example.com,555-867-5309\r\nBob Smith,,555-0100\r\n\r\n"
TXT_TEXT = "Notes from the call\r\n\r\nJane Doe can be reached at jane.doe@example.com\r\n  indented line\r\n"
MD_TEXT = "# Contact\n\n- **Email:** jane.doe@example.com\n- Phone: 555-867-5309\n"
HTML_TEXT = """<!doctype html>
<html><head><title>Secret title</title><style>p { color: red }</style>
<script>var ssn = "999-99-9999";</script></head>
<body>
  <h1>Customer   record</h1>
  <p>Name: Jane&nbsp;Doe &amp; family</p>
  <div>Email: <a href="mailto:x">jane.doe@example.com</a><br>Phone: 555-867-5309</div>
  <table><tr><th>First</th><th>Last</th></tr><tr><td>Jane</td><td>Doe</td></tr></table>
  <noscript>enable javascript</noscript>
  <pre>line one
  line two</pre>
  <span>inline</span> <b>text</b>
</body></html>
"""


def _write_xlsx(path: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "People"
    ws.append(["Name", "Email", "Phone", "Born"])
    ws.append(["Jane Doe", EMAIL, 5558675309.0, dt.date(1990, 5, 17)])
    ws.append(["Bob Smith", None, None, None])
    ws.append([None, None, None, None])  # blank row: no line
    ws["C5"] = "  "                       # whitespace only: no segment
    ws["D6"] = "late value"               # leading empty cells are dropped
    notes = wb.create_sheet("Notes")
    notes["A1"] = "Call Jane on Monday"
    wb.create_sheet("Empty")
    wb.save(path)


def _write_docx(path: Path) -> None:
    import docx

    d = docx.Document()
    d.add_paragraph("Dear Jane Doe,")
    d.add_paragraph("")  # blank: skipped but still counted as paragraph 1
    table = d.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Email"
    table.cell(1, 0).text = "Jane Doe"
    table.cell(1, 1).text = EMAIL
    merged = table.cell(0, 2).merge(table.cell(1, 2))
    merged.text = "merged note"
    d.add_paragraph("Please call 555-867-5309.")
    section = d.sections[0]
    section.header.paragraphs[0].text = "Confidential - Acme HR"
    section.footer.paragraphs[0].text = "Page footer bob@example.org"
    d.save(path)


def _email(body_plain: str | None, body_html: str | None = None) -> EmailMessage:
    msg = EmailMessage()
    msg["From"] = "Jane Doe <jane.doe@example.com>"
    msg["To"] = "Bob Smith <bob@example.org>, carol@example.net"
    msg["Subject"] = "Your account 4111 1111 1111 1111"
    msg["Date"] = "Tue, 06 Oct 2026 10:00:00 +0000"
    if body_plain is not None:
        msg.set_content(body_plain)
        if body_html is not None:
            msg.add_alternative(body_html, subtype="html")
    elif body_html is not None:
        msg.set_content(body_html, subtype="html")
    return msg


def _write_eml(path: Path) -> None:
    msg = _email("Hi Bob,\n\nMy SSN is 123-45-6789.\nThanks\n", "<p>HTML version should not be used</p>")
    msg.add_attachment(b"Attached notes for jane.doe@example.com\n", maintype="text",
                       subtype="plain", filename="notes.txt")
    msg.add_attachment(CSV_TEXT.encode(), maintype="text", subtype="csv", filename="../../people.csv")
    msg.add_attachment(b"\x89PNG\r\n\x1a\n0000", maintype="image", subtype="png", filename="photo.png")
    inner = _email("Forwarded body with 555-867-5309\n")
    inner.add_attachment(b"deep secret 987-65-4321\n", maintype="text", subtype="plain", filename="deep.txt")
    msg.add_attachment(inner)  # message/rfc822
    path.write_bytes(msg.as_bytes())


@pytest.fixture(scope="session")
def samples(tmp_path_factory) -> dict[str, Path]:
    pytest.importorskip("openpyxl")
    pytest.importorskip("docx")
    pytest.importorskip("pypdf")
    root = tmp_path_factory.mktemp("docs")
    paths = {
        "csv": root / "people.csv",
        "tsv": root / "people.tsv",
        "txt": root / "notes.txt",
        "md": root / "contact.md",
        "html": root / "page.html",
        "xlsx": root / "people.xlsx",
        "docx": root / "letter.docx",
        "eml": root / "message.eml",
        "eml_html": root / "html_only.eml",
        "pdf": root / "scan.pdf",
    }
    paths["csv"].write_bytes(CSV_TEXT.encode("utf-8"))
    paths["tsv"].write_text(CSV_TEXT.replace(",", "\t"), encoding="utf-8")
    paths["txt"].write_bytes(("﻿" + TXT_TEXT).encode("utf-8"))
    paths["md"].write_text(MD_TEXT, encoding="utf-8")
    paths["html"].write_text(HTML_TEXT, encoding="utf-8")
    _write_xlsx(paths["xlsx"])
    _write_docx(paths["docx"])
    _write_eml(paths["eml"])
    paths["eml_html"].write_bytes(
        _email(None, "<html><body><p>Hello</p><p>Reach me at jane.doe@example.com</p></body></html>").as_bytes()
    )
    paths["pdf"].write_bytes((DATA / "sample.pdf").read_bytes())
    return paths


def assert_invariants(doc: Document, *, gaps_are_whitespace: bool = True) -> None:
    """Sorted, disjoint, non-empty segments; separators sit between them."""
    assert doc.check() == []
    prev = 0
    for seg in doc.segments:
        piece = doc.text[seg.start:seg.end]
        assert piece.strip(), f"blank segment {seg}"
        assert piece == piece.replace("\r", ""), "newlines must be normalised"
        gap = doc.text[prev:seg.start]
        if gaps_are_whitespace:
            assert gap.strip() == "", f"non-whitespace between segments: {gap!r}"
        prev = seg.end
    if gaps_are_whitespace:
        assert doc.text[prev:].strip() == ""
    assert "\r" not in doc.text


def where_of(doc: Document, value: str, occurrence: int = 0) -> list[dict]:
    i = -1
    for _ in range(occurrence + 1):
        i = doc.text.index(value, i + 1)
    return doc.locate(i, i + len(value))


# ------------------------------------------------------------------ basics


def test_import_is_cheap():
    code = (
        "import sys, opf_eval.documents, opf_eval.fixtures; "
        "bad = [m for m in ('pypdf', 'docx', 'openpyxl') if m in sys.modules]; "
        "print(','.join(bad))"
    )
    out = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True, check=True)
    assert out.stdout.strip() == ""


def test_supported_suffixes_cover_every_kind():
    assert documents.SUPPORTED_SUFFIXES[".xlsm"] == "xlsx"
    assert documents.SUPPORTED_SUFFIXES[".tsv"] == "csv"
    assert documents.SUPPORTED_SUFFIXES[".htm"] == "html"
    assert documents.SUPPORTED_SUFFIXES[".markdown"] == "md"
    assert set(documents.SUPPORTED_SUFFIXES.values()) == {
        "pdf", "docx", "xlsx", "csv", "txt", "md", "html", "eml"
    }


def test_text_builder_and_locate():
    b = TextBuilder()
    b.sep("\n")                 # leading separators are dropped
    assert b.add("Jane", {"line": 1})
    assert not b.add("   ", {"line": 2})
    b.newline()
    b.add("a\r\nb", {"line": 3})
    doc = b.build(id="x", source="s", kind="txt")
    assert doc.text == "Jane\na\nb"
    assert doc.segments == [Segment(0, 4, {"line": 1}), Segment(5, 8, {"line": 3})]
    assert doc.locate(0, 4) == [{"line": 1}]
    assert doc.locate(2, 7) == [{"line": 1}, {"line": 3}]
    assert doc.locate(4, 5) == []           # the separator only
    assert doc.locate(6, 6) == [{"line": 3}]  # zero width -> containing segment
    assert doc.locate(100, 200) == []
    # Returned dicts are copies: mutating one cannot corrupt the document.
    doc.locate(0, 1)[0]["line"] = 99
    assert doc.segments[0].where == {"line": 1}


def test_segment_json_round_trip():
    seg = Segment(3, 9, {"sheet": "Q1", "cell": "B7"})
    assert Segment.from_json(seg.to_json()) == seg
    assert Segment.from_json({"start": 3, "end": 9, "where": {"sheet": "Q1", "cell": "B7"}}) == seg


def test_check_reports_bad_segments():
    doc = Document("x", "s", "txt", "abcdef", [Segment(0, 3, {}), Segment(2, 4, {}), Segment(5, 9, {})])
    problems = doc.check()
    assert any("overlaps" in p for p in problems)
    assert any("out of range" in p for p in problems)


# ----------------------------------------------------------------- parsers


def test_parse_ids_are_content_hash_plus_relative_path(samples):
    p = samples["csv"]
    sha = hashlib.sha1(p.read_bytes()).hexdigest()[:12]
    assert parse(p).id == f"{sha}:people.csv"
    assert parse(p, root=p.parent.parent).id == f"{sha}:{p.parent.name}/people.csv"
    # A root that does not contain the file falls back to the file name.
    assert parse(p, root="/nonexistent").id == f"{sha}:people.csv"
    doc = parse(str(p))
    assert doc.source == str(p) and doc.kind == "csv"


def test_parse_rejects_unsupported_and_missing(tmp_path):
    with pytest.raises(ValueError, match="unsupported file type '.zip'"):
        parse(tmp_path / "a.zip")
    with pytest.raises(FileNotFoundError):
        parse(tmp_path / "missing.txt")


def test_txt(samples):
    doc = parse(samples["txt"])
    assert_invariants(doc)
    assert doc.kind == "txt"
    assert not doc.text.startswith("﻿")
    assert doc.text.splitlines()[0] == "Notes from the call"
    assert where_of(doc, EMAIL) == [{"line": 3}]
    # Lines are kept as written, including indentation.
    assert doc.locate(*_span(doc, "  indented line")) == [{"line": 4}]
    assert [s.where["line"] for s in doc.segments] == [1, 3, 4]


def _span(doc: Document, value: str) -> tuple[int, int]:
    i = doc.text.index(value)
    return i, i + len(value)


def test_markdown_is_kept_as_written(samples):
    doc = parse(samples["md"])
    assert_invariants(doc)
    assert doc.kind == "md"
    assert doc.text == MD_TEXT
    assert where_of(doc, EMAIL) == [{"line": 3}]


def test_csv(samples):
    doc = parse(samples["csv"])
    assert_invariants(doc)
    assert doc.text == "name\temail\tphone\nJane Doe\tjane.doe@example.com\t555-867-5309\nBob Smith\t555-0100"
    assert where_of(doc, EMAIL) == [{"row": 2, "column": "email"}]
    assert where_of(doc, "555-0100") == [{"row": 3, "column": "phone"}]
    assert where_of(doc, "name") == [{"row": 1, "column": "name"}]


def test_tsv_and_sniffed_delimiters(samples, tmp_path):
    doc = parse(samples["tsv"])
    assert_invariants(doc)
    assert where_of(doc, EMAIL) == [{"row": 2, "column": "email"}]

    semi = tmp_path / "semi.csv"
    semi.write_text("id;contact\n1;jane.doe@example.com\n", encoding="utf-8")
    doc = parse(semi)
    assert doc.text == "id\tcontact\n1\tjane.doe@example.com"
    assert where_of(doc, EMAIL) == [{"row": 2, "column": "contact"}]


def test_csv_column_letter_when_header_missing(tmp_path):
    p = tmp_path / "ragged.csv"
    p.write_text('a,,\n1,"two, quoted","multi\r\nline"\n', encoding="utf-8")
    doc = parse(p)
    assert_invariants(doc)
    assert where_of(doc, "two, quoted") == [{"row": 2, "column": "B"}]
    assert where_of(doc, "multi\nline") == [{"row": 2, "column": "C"}]


def test_xlsx(samples):
    doc = parse(samples["xlsx"])
    assert_invariants(doc)
    assert doc.kind == "xlsx"
    assert where_of(doc, EMAIL) == [{"sheet": "People", "cell": "B2"}]
    # Integral floats lose ".0" and dates use ISO format.
    assert where_of(doc, "5558675309") == [{"sheet": "People", "cell": "C2"}]
    assert "5558675309.0" not in doc.text
    assert where_of(doc, "1990-05-17") == [{"sheet": "People", "cell": "D2"}]
    assert where_of(doc, "late value") == [{"sheet": "People", "cell": "D6"}]
    assert where_of(doc, "Call Jane on Monday") == [{"sheet": "Notes", "cell": "A1"}]
    # One line per non-empty row and a blank line between sheets.
    people, notes = doc.text.split("\n\n")
    assert people.splitlines() == [
        "Name\tEmail\tPhone\tBorn",
        f"Jane Doe\t{EMAIL}\t5558675309\t1990-05-17",
        "Bob Smith",
        "late value",
    ]
    assert notes == "Call Jane on Monday"


def test_xlsx_ignores_a_wrong_stored_dimension(samples, tmp_path):
    """Some non-Excel writers store `<dimension ref="A1"/>` for a whole
    sheet. Read-only openpyxl would trust it and drop every other cell."""
    import re
    import zipfile

    bad = tmp_path / "wrongdim.xlsx"
    with zipfile.ZipFile(samples["xlsx"]) as zin, zipfile.ZipFile(bad, "w") as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                data, n = re.subn(rb'<dimension ref="[^"]*"/>', b'<dimension ref="A1"/>', data)
                assert n == 1
            zout.writestr(info, data)
    doc = parse(bad)
    assert doc.text == parse(samples["xlsx"]).text
    assert where_of(doc, EMAIL) == [{"sheet": "People", "cell": "B2"}]
    assert where_of(doc, "late value") == [{"sheet": "People", "cell": "D6"}]


def test_docx(samples):
    import docx

    doc = parse(samples["docx"])
    assert_invariants(doc)
    lines = doc.text.split("\n")
    # Document order: paragraph, table rows, paragraph; then headers/footers.
    assert lines[0] == "Dear Jane Doe,"
    assert lines[1] == "Name\tEmail\tmerged note"
    assert lines[2] == f"Jane Doe\t{EMAIL}"   # the merged cell is emitted once
    assert lines[3] == "Please call 555-867-5309."
    assert doc.text.count("merged note") == 1
    assert where_of(doc, EMAIL) == [{"table": 0, "row": 1, "col": 1}]
    assert where_of(doc, "Dear Jane Doe") == [{"paragraph": 0}]
    # Paragraph indices count the skipped blank paragraph, so they index doc.paragraphs.
    assert where_of(doc, "555-867-5309") == [{"paragraph": 2}]
    assert docx.Document(str(samples["docx"])).paragraphs[2].text == "Please call 555-867-5309."
    assert where_of(doc, "Acme HR") == [{"part": "header", "section": 0}]
    assert where_of(doc, "bob@example.org") == [{"part": "footer", "section": 0}]
    _body, extras = doc.text.split("\n\n")
    assert extras.splitlines() == ["Confidential - Acme HR", "Page footer bob@example.org"]


def _docx_with_hidden_text(path: Path) -> None:
    """Text that python-docx's Paragraph.text and iter_inner_content skip."""
    import docx
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    w = nsdecls("w")
    d = docx.Document()
    d.add_paragraph("intro")
    outro = d.add_paragraph("outro")
    # A body-level content control holding a paragraph and a table.
    outro._p.addprevious(parse_xml(
        f"<w:sdt {w}><w:sdtPr><w:alias w:val='Patient'/></w:sdtPr><w:sdtContent>"
        "<w:p><w:r><w:t>SSN 123-45-6789</w:t></w:r></w:p>"
        "<w:tbl><w:tr><w:tc><w:p><w:r><w:t>Jane</w:t></w:r></w:p></w:tc>"
        "<w:tc><w:p><w:r><w:t>Doe</w:t></w:r></w:p></w:tc></w:tr></w:tbl>"
        "</w:sdtContent></w:sdt>"
    ))
    p = d.add_paragraph("Contact: ")  # tracked insertion
    p._p.append(parse_xml(f"<w:ins {w} w:id='1' w:author='a'><w:r><w:t>{EMAIL}</w:t></w:r></w:ins>"))
    p = d.add_paragraph("Name: ")  # inline content control
    p._p.append(parse_xml(f"<w:sdt {w}><w:sdtContent><w:r><w:t>Jane Doe</w:t></w:r></w:sdtContent></w:sdt>"))
    p = d.add_paragraph("Ref: ")  # simple field
    p._p.append(parse_xml(f"<w:fldSimple {w} w:instr='REF x'><w:r><w:t>555-867-5309</w:t></w:r></w:fldSimple>"))
    p = d.add_paragraph("Kept")  # tracked deletion and a tab
    p._p.append(parse_xml(f"<w:del {w} w:id='2' w:author='a'><w:r><w:delText>deleted 999-99-9999</w:delText></w:r></w:del>"))
    p._p.append(parse_xml(f"<w:r {w}><w:tab/><w:t>after tab</w:t></w:r>"))
    table = d.add_table(rows=1, cols=1)  # content control inside a table cell
    table.cell(0, 0)._tc.append(parse_xml(
        f"<w:sdt {w}><w:sdtContent><w:p><w:r><w:t>cell control 4111</w:t></w:r></w:p></w:sdtContent></w:sdt>"
    ))
    d.save(path)


def test_docx_reads_content_controls_and_tracked_changes(tmp_path):
    pytest.importorskip("docx")
    path = tmp_path / "forms.docx"
    _docx_with_hidden_text(path)
    doc = parse(path)
    assert_invariants(doc)
    assert doc.warnings == []
    assert doc.text.split("\n") == [
        "intro",
        "SSN 123-45-6789",
        "Jane\tDoe",
        "outro",
        f"Contact: {EMAIL}",
        "Name: Jane Doe",
        "Ref: 555-867-5309",
        "Kept\tafter tab",
        "cell control 4111",
    ]
    assert "999-99-9999" not in doc.text  # deleted text stays out
    assert where_of(doc, "123-45-6789") == [{"content_control": 0, "paragraph": 0}]
    assert where_of(doc, "Doe") == [{"content_control": 0, "table": 0, "row": 0, "col": 1}]
    # Body paragraph numbering still matches doc.paragraphs.
    import docx

    paragraphs = docx.Document(str(path)).paragraphs
    assert where_of(doc, "outro") == [{"paragraph": 1}]
    assert paragraphs[1].text == "outro"
    assert where_of(doc, EMAIL) == [{"paragraph": 2}]
    assert where_of(doc, "cell control") == [{"table": 0, "row": 0, "col": 0}]


def test_docx_warns_about_text_it_cannot_place(tmp_path):
    pytest.importorskip("docx")
    import docx
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    d = docx.Document()
    d.add_paragraph("visible")
    table = d.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "plain cell"
    # A table row wrapped in a content control is not a row python-docx sees.
    table._tbl.append(parse_xml(
        f"<w:sdt {nsdecls('w')}><w:sdtContent><w:tr><w:tc><w:p><w:r><w:t>hidden row</w:t></w:r></w:p>"
        "</w:tc></w:tr></w:sdtContent></w:sdt>"
    ))
    path = tmp_path / "rowsdt.docx"
    d.save(path)
    doc = parse(path)
    assert "visible" in doc.text and "plain cell" in doc.text
    assert len(doc.warnings) == 1
    assert doc.warnings[0].startswith("9 characters of body text could not be extracted")


def test_pdf_with_a_password_is_skipped_with_a_clear_reason(samples, tmp_path):
    from pypdf import PdfWriter

    locked = tmp_path / "locked.pdf"
    w = PdfWriter(clone_from=str(samples["pdf"]))
    w.encrypt("secret", algorithm="RC4-128")
    w.write(locked)
    open_ = tmp_path / "permissions_only.pdf"
    w = PdfWriter(clone_from=str(samples["pdf"]))
    w.encrypt("", owner_password="owner", algorithm="RC4-128")
    w.write(open_)

    docs, skipped = parse_many([locked, open_])
    assert skipped == [{"path": str(locked), "reason": "ValueError: PDF is encrypted and needs a password"}]
    assert [d.source for d in docs] == [str(open_)]
    assert where_of(docs[0], EMAIL) == [{"page": 1}]


def test_pdf(samples):
    doc = parse(samples["pdf"])
    assert_invariants(doc)
    assert doc.kind == "pdf"
    assert where_of(doc, EMAIL) == [{"page": 1}]
    assert where_of(doc, "123-45-6789") == [{"page": 2}]
    assert [s.where for s in doc.segments] == [{"page": 1}, {"page": 2}]
    assert doc.warnings == ["page 3 has no text layer (needs OCR)"]
    assert "\n\n" in doc.text  # pages are joined by a blank line
    assert (DATA / "sample.pdf").stat().st_size < 5_000


def test_html(samples):
    doc = parse(samples["html"])
    assert_invariants(doc)
    lines = doc.text.split("\n")
    assert lines == [
        "Customer record",
        "Name: Jane Doe & family",  # &nbsp; collapses like other whitespace
        f"Email: {EMAIL}",
        "Phone: 555-867-5309",
        "First\tLast",
        "Jane\tDoe",
        "line one",
        "line two",
        "inline text",
    ]
    for hidden in ("Secret title", "999-99-9999", "color: red", "enable javascript"):
        assert hidden not in doc.text
    assert where_of(doc, EMAIL) == [{"line": 3}]


def test_html_edge_cases():
    # A missing </head> must not hide the body; unclosed skip tags end at their parent.
    assert html_to_lines("<html><head><title>t</title><body><p>visible</p>") == ["visible"]
    assert html_to_lines("<p>a<br/>b</p><ul><li>one</li><li>two</li></ul>") == ["a", "b", "one", "two"]
    assert html_to_lines("") == []
    assert html_to_lines("<script>x</script>  ") == []
    assert html_to_lines("<p>&lt;tag&gt; &#169;</p>") == ["<tag> ©"]


def test_html_title_and_unclosed_head_follow_browser_rules():
    # <head> is optional, so a bare <title> is still head content.
    assert html_to_lines("<!doctype html><title>Secret Title Jane</title><p>body</p>") == ["body"]
    # Without </head> or <body>, the first body-only tag ends the head.
    page = "<html><head><meta charset=utf-8><title>T</title><p>Jane Doe jane@example.com</p></html>"
    assert html_to_lines(page) == ["Jane Doe jane@example.com"]
    # Tags allowed in a head keep it open, and tags inside a title do not end it.
    page = "<head><title>a <b>x</b></title><link rel=x><noscript><style>s</style></noscript><p>shown"
    assert html_to_lines(page) == ["shown"]


def test_text_encodings_are_detected_or_reported(tmp_path):
    cases = {
        "bom16.txt": ("\ufeffJane Doe jane@example.com".encode("utf-16-le"), []),
        "bom16be.txt": ("\ufeffJane Doe jane@example.com".encode("utf-16-be"), []),
        "bom32.txt": ("\ufeffJane Doe jane@example.com".encode("utf-32-le"), []),
        "nobom16.txt": ("Jane Doe jane@example.com".encode("utf-16-le"),
                        ["no byte-order mark; decoded as UTF-16-LE"]),
    }
    for name, (data, warnings) in cases.items():
        (tmp_path / name).write_bytes(data)
        doc = parse(tmp_path / name)
        assert doc.text == "Jane Doe jane@example.com", name
        assert doc.warnings == warnings, name

    cp = tmp_path / "excel.csv"
    cp.write_bytes("name,city\nJosé Núñez,Málaga\n".encode("cp1252"))
    doc = parse(cp)
    assert where_of(doc, "José Núñez") == [{"row": 2, "column": "name"}]
    assert doc.warnings == ["text is not valid UTF-8; decoded as Windows-1252"]

    page = tmp_path / "page.html"
    page.write_bytes("<p>José</p>".encode("cp1252"))
    assert parse(page).text == "José"
    assert parse(page).warnings == ["text is not valid UTF-8; decoded as Windows-1252"]

    # Bytes Windows-1252 cannot decode either are replaced and counted.
    bad = tmp_path / "bad.txt"
    bad.write_bytes(b"ok \x81\x8d end")
    doc = parse(bad)
    assert doc.text == "ok \ufffd\ufffd end"
    assert doc.warnings == ["text is not valid UTF-8; 2 undecodable byte sequence(s) replaced with U+FFFD"]


def test_empty_result_is_a_warning(tmp_path):
    p = tmp_path / "blank.txt"
    p.write_text("\n  \n", encoding="utf-8")
    assert parse(p).warnings == ["no text was extracted"]


def test_eml(samples):
    doc = parse(samples["eml"])
    assert_invariants(doc, gaps_are_whitespace=False)
    assert doc.kind == "eml"
    head = doc.text.split("\n\n")[0].split("\n")
    assert head[0] == "From: Jane Doe <jane.doe@example.com>"
    assert head[1].startswith("To: Bob Smith <bob@example.org>")
    assert [line.split(":")[0] for line in head] == ["From", "To", "Subject", "Date"]
    # Header segments cover the value only.
    first = doc.segments[0]
    assert doc.text[first.start:first.end] == "Jane Doe <jane.doe@example.com>"
    assert where_of(doc, EMAIL) == [{"part": "header", "name": "From"}]
    assert where_of(doc, "4111 1111 1111 1111") == [{"part": "header", "name": "Subject"}]
    # Plain body is preferred over the HTML alternative; lines count blanks.
    assert "HTML version" not in doc.text
    assert where_of(doc, "123-45-6789") == [{"part": "body", "line": 3}]
    # Attachments: parsed and located under their (sanitised) filename.
    assert where_of(doc, EMAIL, occurrence=1) == [{"attachment": "notes.txt", "line": 1}]
    assert where_of(doc, "555-0100") == [{"attachment": "people.csv", "row": 3, "column": "phone"}]
    assert where_of(doc, "Forwarded body") == [{"attachment": "attached.eml", "part": "body", "line": 1}]
    assert where_of(doc, "987-65-4321") == [{"attachment": "attached.eml/deep.txt", "line": 1}]
    assert any("photo.png" in w and "unsupported" in w for w in doc.warnings)
    # Every gap between segments is whitespace or a header label.
    prev = 0
    for seg in doc.segments:
        gap = doc.text[prev:seg.start].strip()
        assert gap == "" or gap.rstrip(":") in {"From", "To", "Cc", "Bcc", "Reply-To", "Subject", "Date"}
        prev = seg.end


def test_eml_html_body_and_no_body(samples, tmp_path):
    doc = parse(samples["eml_html"])
    assert_invariants(doc, gaps_are_whitespace=False)
    assert where_of(doc, EMAIL, occurrence=1) == [{"part": "body", "line": 2}]
    assert "<p>" not in doc.text

    bare = tmp_path / "bare.eml"
    bare.write_bytes(b"From: a@example.com\nSubject: hi\n\n")
    doc = parse(bare)
    assert doc.text == "From: a@example.com\nSubject: hi"
    assert any("body" in w for w in doc.warnings)


def test_eml_body_with_unknown_charset_is_decoded_leniently(tmp_path):
    p = tmp_path / "odd.eml"
    p.write_bytes(
        b"From: a@example.com\r\nSubject: x\r\n"
        b"Content-Type: text/plain; charset=x-unknown\r\nContent-Transfer-Encoding: 8bit\r\n\r\n"
        b"Jos\xe9 N\xfa\xf1ez\r\njane.doe@example.com\r\n"
    )
    doc = parse(p)
    assert where_of(doc, "José Núñez") == [{"part": "body", "line": 1}]
    assert where_of(doc, EMAIL) == [{"part": "body", "line": 2}]
    assert "\r" not in doc.text
    assert len(doc.warnings) == 1
    assert "x-unknown" in doc.warnings[0] and "Windows-1252" in doc.warnings[0]


def test_eml_broken_attachment_is_a_warning(tmp_path):
    msg = _email("body\n")
    msg.add_attachment(b"not really a docx", maintype="application",
                       subtype="vnd.openxmlformats-officedocument.wordprocessingml.document",
                       filename="broken.docx")
    p = tmp_path / "broken.eml"
    p.write_bytes(msg.as_bytes())
    doc = parse(p)
    assert where_of(doc, "body") == [{"part": "body", "line": 1}]
    assert any("broken.docx" in w and "could not be parsed" in w for w in doc.warnings)


def test_from_text():
    doc = from_text("Jane Doe\r\n\r\nemail jane.doe@example.com")
    assert_invariants(doc)
    assert doc.kind == "text" and doc.source == "<pasted>"
    assert doc.text == "Jane Doe\n\nemail jane.doe@example.com"
    assert where_of(doc, EMAIL) == [{"line": 3}]
    assert doc.id == from_text("Jane Doe\n\nemail jane.doe@example.com").id  # stable
    assert doc.id.endswith(":pasted")
    assert from_text("x", id="mine", source="ticket 7").id == "mine"
    assert from_text("").segments == []


# --------------------------------------------------------------- parse_many


def test_parse_many_walks_folders_and_skips(samples, tmp_path):
    root = tmp_path / "inbox"
    (root / "sub").mkdir(parents=True)
    (root / ".hidden").mkdir()
    (root / "a.txt").write_text("alpha jane.doe@example.com\n")
    (root / "sub" / "a.txt").write_text("alpha in a subfolder\n")
    (root / "sub" / "people.xlsx").write_bytes(samples["xlsx"].read_bytes())
    (root / ".secret.txt").write_text("hidden file\n")
    (root / ".hidden" / "x.txt").write_text("hidden folder\n")
    (root / "archive.zip").write_bytes(b"PK")
    (root / "~$letter.docx").write_bytes(b"lock")
    (root / "broken.docx").write_bytes(b"this is not a zip file")

    docs, skipped = parse_many(root)
    ids = [d.id.split(":", 1)[1] for d in docs]
    assert ids == ["a.txt", "sub/a.txt", "sub/people.xlsx"]
    reasons = {Path(s["path"]).name: s["reason"] for s in skipped}
    assert set(reasons) == {"archive.zip", "~$letter.docx", "broken.docx"}
    assert "unsupported" in reasons["archive.zip"]
    assert "lock" in reasons["~$letter.docx"]
    assert reasons["broken.docx"].startswith(("BadZipFile", "PackageNotFoundError", "KeyError", "ValueError"))

    shallow, _ = parse_many(root, recursive=False)
    assert [d.id.split(":", 1)[1] for d in shallow] == ["a.txt"]


def test_parse_many_lists_and_duplicates(samples, tmp_path):
    docs, skipped = parse_many([samples["csv"], str(samples["csv"]), tmp_path / "nope.txt", samples["md"]])
    assert [d.kind for d in docs] == ["csv", "md"]
    assert {s["reason"] for s in skipped} == {"not found", f"duplicate of {samples['csv']}"}
    docs, skipped = parse_many(samples["txt"])
    assert len(docs) == 1 and skipped == []


# ----------------------------------------------------------------- chunking


def assert_chunks(doc: Document, ranges: list[tuple[int, int]], max_chars: int) -> None:
    assert ranges[0][0] == 0 and ranges[-1][1] == len(doc.text)
    for (_, b), (c, _) in itertools.pairwise(ranges):
        assert b == c, "chunks must not leave gaps or overlap"
    for a, b in ranges:
        assert 0 < b - a <= max_chars


def test_chunk_cuts_only_at_segment_boundaries(samples):
    doc = parse(samples["xlsx"])
    for max_chars in (15, 30, 64, 10_000):
        ranges = chunk(doc, max_chars=max_chars)
        assert_chunks(doc, ranges, max_chars)
        cuts = {b for _, b in ranges[:-1]}
        for seg in doc.segments:
            if seg.end - seg.start <= max_chars:
                assert not any(seg.start < c < seg.end for c in cuts), (max_chars, seg)
    assert chunk(doc, max_chars=10_000) == [(0, len(doc.text))]


def test_chunk_splits_one_long_segment_softly():
    words = " ".join(f"word{i:03d}." for i in range(200))  # one 1,799-char line
    doc = from_text(f"short\n{words}\nend")
    ranges = chunk(doc, max_chars=100)
    assert_chunks(doc, ranges, 100)
    # Cuts inside the long line land just after ". " (never mid-word).
    for _, b in ranges[:-1]:
        assert doc.text[b - 1] in " \n"
    # The short first line ends its own chunk because the next segment does not fit.
    assert ranges[0] == (0, 6)


def test_chunk_hard_cut_and_empty():
    doc = from_text("x" * 25)
    assert chunk(doc, max_chars=10) == [(0, 10), (10, 20), (20, 25)]
    assert chunk(from_text(""), max_chars=10) == []
    with pytest.raises(ValueError):
        chunk(doc, max_chars=0)


def test_chunk_never_emits_a_separator_only_chunk():
    # The cut after "a"*9 used to leave the "\n\n" separator as its own chunk.
    doc = from_text("a" * 9 + "\n\n" + "b" * 30)
    ranges = chunk(doc, max_chars=10)
    assert_chunks(doc, ranges, 10)
    assert [doc.text[a:b] for a, b in ranges] == ["a" * 9, "\n\n" + "b" * 8, "b" * 10, "b" * 10, "b" * 2]
    # Trailing whitespace joins the last chunk when it fits.
    doc = from_text("abc\n\n\n")
    assert chunk(doc, max_chars=10) == [(0, 6)]


def test_chunk_fuzz_has_no_blank_chunks_when_avoidable():
    import random

    rng = random.Random(13)
    for _ in range(2000):
        lines = ["x" * rng.randint(0, 25) for _ in range(rng.randint(1, 8))]
        doc = from_text(rng.choice(["\n", "\n\n", " \n"]).join(lines))
        if not doc.text:
            continue
        max_chars = rng.randint(3, 30)
        ranges = chunk(doc, max_chars=max_chars)
        assert_chunks(doc, ranges, max_chars)
        for k, (a, b) in enumerate(ranges):
            piece = doc.text[a:b]
            if not piece.strip():
                # Only allowed when no neighbour could absorb it.
                prev_full = k == 0 or b - ranges[k - 1][0] > max_chars
                next_full = k == len(ranges) - 1 or ranges[k + 1][1] - a > max_chars
                assert b - a > 0 and prev_full and next_full, (doc.text, max_chars, ranges)
        # A line that fits in a chunk is never cut.
        cuts = {b for _, b in ranges[:-1]}
        for seg in doc.segments:
            if seg.end - seg.start <= max_chars:
                assert not any(seg.start < c < seg.end for c in cuts), (doc.text, max_chars, ranges, seg)


def test_chunk_keeps_a_segment_whole_when_separators_push_it_past_the_limit():
    # After a cut at a segment end the next chunk starts on the "\n\n"
    # separator. Separator and page together do not fit, but the page alone
    # does, so the page must not be cut through "Jane Doe 123".
    b = TextBuilder()
    b.add("x" * 12, {"page": 1})
    b.blank_line()
    b.add("Jane Doe 123", {"page": 2})
    doc = b.build(id="t", source="t", kind="pdf")
    ranges = chunk(doc, max_chars=12)
    assert_chunks(doc, ranges, 12)
    assert [doc.text[a:b] for a, b in ranges] == ["x" * 12, "\n\n", "Jane Doe 123"]
    # The separators join the previous chunk when there is room for them.
    b = TextBuilder()
    b.add("x" * 8, {"page": 1})
    b.blank_line()
    b.add("Jane Doe 123", {"page": 2})
    doc = b.build(id="t", source="t", kind="pdf")
    assert [doc.text[a:b] for a, b in chunk(doc, max_chars=12)] == ["x" * 8 + "\n\n", "Jane Doe 123"]


def test_chunk_text_without_segments():
    doc = Document("x", "s", "text", "aaa bbb ccc ddd", [])
    ranges = chunk(doc, max_chars=8)
    assert_chunks(doc, ranges, 8)
    assert [doc.text[a:b] for a, b in ranges] == ["aaa bbb ", "ccc ddd"]
